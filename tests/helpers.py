"""Builders for synthetic fixtures. No real client documents anywhere in tests."""

from __future__ import annotations

import base64
import email.utils
import io
import zipfile
from datetime import datetime, timezone
from email.message import EmailMessage

import openpyxl
from openpyxl.workbook.defined_name import DefinedName

# 1x1 transparent PNG.
PNG_1X1 = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNkYPhfDwAChwGA60e6kgAAAABJRU5ErkJggg=="
)


def _pdf(objects: list[bytes]) -> bytes:
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return bytes(out)


def make_pdf(pages: int = 1, text: str | None = "Invoice 12345 Vendor ACME Total 1000.00", acroform: bool = False) -> bytes:
    """Tiny valid PDF. ``text=None`` gives pages with no text layer (a 'scan')."""
    page_ids = [4 + 2 * i for i in range(pages)]
    kids = " ".join(f"{pid} 0 R" for pid in page_ids)
    catalog = f"<< /Type /Catalog /Pages 2 0 R{' /AcroForm << /Fields [] >>' if acroform else ''} >>".encode()
    pages_obj = f"<< /Type /Pages /Kids [{kids}] /Count {pages} >>".encode()
    font = b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"
    objects = [catalog, pages_obj, font]
    for i in range(pages):
        content_id = page_ids[i] + 1
        page = (
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {content_id} 0 R >>"
        ).encode()
        stream = f"BT /F1 12 Tf 72 720 Td ({text} page {i + 1}) Tj ET".encode() if text else b""
        content = b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream"
        objects += [page, content]
    return _pdf(objects)


def make_xlsx() -> bytes:
    wb = openpyxl.Workbook()
    terms = wb.active
    terms.title = "Terms"
    terms["A1"], terms["B1"] = "Vendor", "ACME"
    terms["A2"], terms["B2"] = "Allowance", 0.025
    terms["A4"] = "Promotional terms"
    terms.merge_cells("A4:C4")
    calc = wb.create_sheet("Calc")
    calc["A1"] = "=Terms!B2*100"
    calc["A2"] = "=SUM(A1:A1)"
    calc["A3"] = 42
    lookup = wb.create_sheet("Lookup")
    lookup.sheet_state = "hidden"
    lookup["A1"] = "x"
    wb.defined_names["AllowanceRate"] = DefinedName("AllowanceRate", attr_text="Terms!$B$2")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def make_csv(delimiter: str = ",", rows: int = 5, encoding: str = "utf-8") -> bytes:
    header = delimiter.join(["invoice_no", "vendor", "amount"])
    lines = [header] + [delimiter.join([f"INV{i:04d}", "ACMÉ Ltd", f"{i * 10.5:.2f}"]) for i in range(1, rows + 1)]
    return "\n".join(lines).encode(encoding)


def make_zip(members: dict[str, bytes]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, data in members.items():
            zf.writestr(name, data)
    return buf.getvalue()


def make_email(
    *,
    subject: str = "RE: Invoice INV0001",
    sender: str = "AP Team <ap@retailer.example>",
    attachments: dict[str, bytes] | None = None,
    nested: list[tuple[EmailMessage, str | None]] | None = None,
    inline_png: bool = False,
    html: bool = True,
    when: datetime | None = None,
) -> EmailMessage:
    msg = EmailMessage()
    msg["From"] = sender
    msg["To"] = "claims@vendor.example"
    msg["Subject"] = subject
    msg["Date"] = email.utils.format_datetime(when or datetime(2024, 3, 5, 10, 0, tzinfo=timezone.utc))
    msg["Message-ID"] = email.utils.make_msgid(domain="retailer.example")
    msg.set_content("Please see attached.")
    if html:
        msg.add_alternative("<html><body><p>Please see attached.</p></body></html>", subtype="html")
    for name, data in (attachments or {}).items():
        maintype, subtype = _mime_for(name)
        msg.add_attachment(data, maintype=maintype, subtype=subtype, filename=name)
    if inline_png:
        msg.add_attachment(PNG_1X1, maintype="image", subtype="png", filename="logo.png", disposition="inline", cid="<logo@retailer.example>")
    for inner, filename in nested or []:
        if filename:
            msg.add_attachment(inner, filename=filename)
        else:
            msg.add_attachment(inner)
    return msg


def _mime_for(name: str) -> tuple[str, str]:
    ext = name.rsplit(".", 1)[-1].lower()
    return {
        "pdf": ("application", "pdf"),
        "xlsx": ("application", "vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
        "csv": ("text", "csv"),
        "zip": ("application", "zip"),
        "png": ("image", "png"),
        "txt": ("text", "plain"),
    }.get(ext, ("application", "octet-stream"))
