"""Spreadsheet probes for the census.

What the census needs to know about a workbook is how hard it will be
later: how many sheets, whether formulas cross sheets or reach into other
workbooks, hidden sheets, macros, pivot caches, merged ranges (a strong
hint that the sheet is a form rather than a table). Values are not read
into anything; that is the spreadsheet path's job in a later phase.
"""

from __future__ import annotations

import io
import re
import zipfile

import openpyxl
import pyxlsb

from docintel.models import ArtifactRecord

READ_ONLY_ABOVE_BYTES = 15 * 1024 * 1024
MAX_CELLS = 2_000_000

# Sheet1!A1, 'Rate Card'!B2:B9, and [1]Sheet1!A1 for external workbooks.
_SHEET_REF = re.compile(r"(?:'[^']{1,255}'|[A-Za-z_][A-Za-z0-9_.]{0,254})!")
_EXTERNAL_REF = re.compile(r"\[\d+\]")


def _zip_level(record: ArtifactRecord, data: bytes) -> None:
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        names = zf.namelist()
    record.external_link_count = sum(1 for n in names if n.startswith("xl/externalLinks/externalLink"))
    record.pivot_cache_count = sum(1 for n in names if n.startswith("xl/pivotCache/pivotCacheDefinition"))
    record.has_vba = "xl/vbaProject.bin" in names


def _dims(ws) -> tuple[int, int]:
    try:
        return int(ws.max_row or 0), int(ws.max_column or 0)
    except Exception:  # noqa: BLE001 - unsized read-only sheets
        return 0, 0


def formula_text(value) -> str | None:
    if isinstance(value, str):
        return value if value.startswith("=") else None
    text = getattr(value, "text", None)  # ArrayFormula, DataTableFormula
    if isinstance(text, str):
        return text if text.startswith("=") else "=" + text
    return None


def probe_xlsx(record: ArtifactRecord, data: bytes) -> None:
    _zip_level(record, data)
    read_only = len(data) > READ_ONLY_ABOVE_BYTES
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=read_only, data_only=False, keep_links=True)
    try:
        sheets = wb.worksheets
        record.sheet_count = len(wb.sheetnames)
        record.hidden_sheet_count = sum(1 for ws in sheets if getattr(ws, "sheet_state", "visible") != "visible")
        defined = len(wb.defined_names)
        for ws in sheets:
            defined += len(getattr(ws, "defined_names", {}) or {})
        record.defined_name_count = defined

        max_rows = max_cols = nonempty = formulas = cross = merged = tables = 0
        budget = MAX_CELLS
        truncated = False
        for ws in sheets:
            if not read_only:
                merged += len(ws.merged_cells.ranges)
                tables += len(getattr(ws, "tables", {}) or {})
            rows, cols = _dims(ws)
            max_rows, max_cols = max(max_rows, rows), max(max_cols, cols)
            if budget <= 0:
                truncated = True
                continue
            for row in ws.iter_rows(values_only=True):
                budget -= len(row)
                for value in row:
                    if value is None or value == "":
                        continue
                    nonempty += 1
                    formula = formula_text(value)
                    if formula is not None:
                        formulas += 1
                        if _SHEET_REF.search(formula) or _EXTERNAL_REF.search(formula):
                            cross += 1
                if budget <= 0:
                    truncated = True
                    break

        record.max_rows, record.max_cols = max_rows, max_cols
        record.nonempty_cells = nonempty
        record.formula_count = formulas
        record.cross_sheet_formula_count = cross
        record.merged_range_count = None if read_only else merged
        record.table_count = None if read_only else tables
        record.probe_truncated = True if truncated else None
    finally:
        wb.close()


def probe_xlsb(record: ArtifactRecord, data: bytes) -> None:
    """Binary workbooks: pyxlsb exposes values but not formulas, so formula
    fields stay None and the report says so."""
    _zip_level(record, data)
    wb = pyxlsb.open_workbook(io.BytesIO(data))
    try:
        record.sheet_count = len(wb.sheets)
        max_rows = max_cols = nonempty = 0
        budget = MAX_CELLS
        truncated = False
        for name in wb.sheets:
            if budget <= 0:
                truncated = True
                break
            with wb.get_sheet(name) as sheet:
                for row in sheet.rows(sparse=True):
                    budget -= len(row)
                    for cell in row:
                        if cell.v is None or cell.v == "":
                            continue
                        nonempty += 1
                        max_rows = max(max_rows, cell.r + 1)
                        max_cols = max(max_cols, cell.c + 1)
                    if budget <= 0:
                        truncated = True
                        break
        record.max_rows, record.max_cols = max_rows, max_cols
        record.nonempty_cells = nonempty
        record.probe_truncated = True if truncated else None
    finally:
        close = getattr(wb, "close", None)
        if close is not None:
            close()
